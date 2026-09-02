from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


HEADER_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def write_json(payload: dict[str, Any], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        json.dump(payload, handle, ensure_ascii=False, indent=2)


def _metric_frame(payload: dict[str, Any]) -> pd.DataFrame:
    return pd.DataFrame({"metric": list(payload.keys()), "value": list(payload.values())})


def _format_workbook(writer: pd.ExcelWriter) -> None:
    workbook = writer.book
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column_cells in worksheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column_cells[:250]]
            width = min(max(max((len(value) for value in values), default=8) + 2, 11), 42)
            worksheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def write_excel_workbench(
    output_path: Path,
    portfolio_kpis: dict[str, Any],
    segment_table: pd.DataFrame,
    pricing_metrics: dict[str, Any],
    pricing_scores: pd.DataFrame,
    reserve_metrics: dict[str, Any],
    reserve_by_period: pd.DataFrame,
    development_factors: pd.DataFrame,
    fraud_metrics: dict[str, Any],
    fraud_scores: pd.DataFrame,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        _metric_frame(portfolio_kpis).to_excel(writer, sheet_name="Portfolio KPI", index=False)
        segment_table.to_excel(writer, sheet_name="Segment Performance", index=False)
        _metric_frame(pricing_metrics).to_excel(writer, sheet_name="Pricing Validation", index=False)
        (
            pricing_scores.sort_values("pricing_adequacy_index")
            .head(2_000)
            .to_excel(writer, sheet_name="Pricing Review", index=False)
        )
        _metric_frame(reserve_metrics).to_excel(writer, sheet_name="Reserve KPI", index=False)
        reserve_by_period.to_excel(writer, sheet_name="Reserve by Quarter", index=False)
        development_factors.to_excel(writer, sheet_name="Development Factors", index=False)
        _metric_frame(fraud_metrics).to_excel(writer, sheet_name="Triage KPI", index=False)
        (
            fraud_scores.loc[fraud_scores["review_recommended"]]
            .sort_values(["triage_score", "incurred_amount"], ascending=False)
            .head(2_000)
            .to_excel(writer, sheet_name="Human Review Queue", index=False)
        )
        _format_workbook(writer)


def write_executive_summary(
    output_path: Path,
    portfolio: dict[str, Any],
    pricing: dict[str, Any],
    reserve: dict[str, Any],
    fraud: dict[str, Any],
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    content = f"""# Executive Analytical Summary

## Portfolio position

- Policies: {portfolio['policy_count']:,}
- Exposure years: {portfolio['exposure_years']:,.2f}
- Earned premium: TRY {portfolio['earned_premium']:,.2f}
- Incurred claims: TRY {portfolio['incurred_claims']:,.2f}
- Claim frequency: {portfolio['claim_frequency']:.2%}
- Average incurred severity: TRY {portfolio['average_incurred_severity']:,.2f}
- Loss ratio: {portfolio['loss_ratio']:.2%}
- Combined ratio: {portfolio['combined_ratio']:.2%}

## Technical pricing validation

- OOT actual frequency: {pricing['oot_actual_frequency']:.2%}
- OOT predicted frequency: {pricing['oot_predicted_frequency']:.2%}
- OOT frequency calibration ratio: {pricing['oot_frequency_calibration_ratio']:.3f}
- OOT pure-premium calibration ratio: {pricing['oot_pure_premium_calibration_ratio']:.3f}
- Policies in indicated increase review: {pricing['underpriced_policy_share']:.2%}

## Reserving

- Chain Ladder IBNR: TRY {reserve['chain_ladder_ibnr']:,.2f}
- Bornhuetter-Ferguson IBNR: TRY {reserve['bornhuetter_ferguson_ibnr']:,.2f}
- Chain Ladder total error vs synthetic truth: {reserve['chain_ladder_total_error_pct']:.2%}
- Bornhuetter-Ferguson total error vs synthetic truth: {reserve['bornhuetter_ferguson_total_error_pct']:.2%}

## Human-controlled fraud triage

- Review queue size: {fraud['alert_count']:,} claims ({fraud['alert_rate']:.2%})
- Precision at review rate: {fraud['precision_at_alert_rate']:.2%}
- Lift versus random review: {fraud['precision_lift_vs_random']:.2f}x
- Automated claim denials: {fraud['automated_claim_denials']}

## Governance boundary

All customer, policy, claim, pricing, reserve, and fraud labels are synthetic. The
models are portfolio demonstrations only. They are not approved for real pricing,
underwriting, reserving, claim settlement, or adverse-action decisions. Authorized
actuarial, legal, compliance, model-risk, and human review is required before any
real-world use.
"""
    output_path.write_text(content, encoding="utf-8")

